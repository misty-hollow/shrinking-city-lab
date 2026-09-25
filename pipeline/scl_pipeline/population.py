"""공주시 주민등록 인구(2026-08-31) 원본 xlsx 읽기.

- 법정동별통반별인구현황: 읍·면 시트마다 `법정동명 | 통 | 반 | 세대수 | 총인구수 ...` 행.
  법정리 인구 = 같은 법정동명의 통·반 행 총인구수 합(설계안 4-1).
- 인구 및 세대현황: 읍·면 총인구(시트 `인구 및 세대현황`), 65세 이상(시트
  `기관별연령별인구통계(65세이상)` 합계 행).
"""

from __future__ import annotations

import collections
from dataclasses import dataclass
from pathlib import Path

import openpyxl

from . import config


@dataclass(frozen=True)
class Population:
    reference_date: str
    # 읍·면 → [(법정리, 인구)] (원본에 처음 나온 순서)
    ri: dict[str, list[tuple[str, int]]]
    emd_total: dict[str, int]
    emd_65plus: dict[str, int]
    city_total: int
    city_65plus: int
    dong_total: dict[str, int]


def _int(v: object) -> int:
    return int(str(v).replace(",", "").strip())


def read_population(tongban_xlsx: Path, households_xlsx: Path) -> Population:
    wb = openpyxl.load_workbook(tongban_xlsx, read_only=True, data_only=True)
    ri: dict[str, list[tuple[str, int]]] = {}
    ref_dates = set()
    for em in config.EMD_NAMES:
        ws = wb[em]
        agg: dict[str, int] = collections.OrderedDict()
        total = None
        in_table = False
        for r in ws.iter_rows(values_only=True):
            if r[0] == "작성일자":
                ref_dates.add(str(r[1]).strip())
            if r[0] == "총인구수":
                total = _int(r[1])
            if r[0] == "법정동명":
                in_table = True
                continue
            if in_table and r[0] and isinstance(r[4], (int, float)):
                agg[r[0]] = agg.get(r[0], 0) + int(r[4])
        if total is None or sum(agg.values()) != total:
            raise ValueError(f"{em}: 통·반 합 {sum(agg.values())} != 총인구수 {total}")
        ri[em] = list(agg.items())
    wb.close()
    if len(ref_dates) != 1:
        raise ValueError(f"작성일자가 시트마다 다르다: {ref_dates}")

    wb1 = openpyxl.load_workbook(households_xlsx, read_only=True, data_only=True)
    rows = list(wb1["기관별연령별인구통계(65세이상)"].iter_rows(values_only=True))
    header, totals = rows[5], rows[7]
    if str(totals[1]).replace(" ", "") != "합계":
        raise ValueError("65세 이상 시트의 합계 행 위치가 바뀌었다")
    old = {h: _int(totals[j]) for j, h in enumerate(header) if h and j >= 3}
    rows = list(wb1["인구 및 세대현황"].iter_rows(values_only=True))
    allp: dict[str, int] = {}
    for r in rows:
        for j, c in enumerate(r or ()):
            if isinstance(c, str) and c.strip() in (*config.EMD_NAMES, *config.DONGS, "공주시"):
                name = c.strip()
                if name not in allp:
                    allp[name] = _int(next(x for x in r[j + 1 :] if x is not None))
    wb1.close()

    for em in config.EMD_NAMES:
        if sum(n for _, n in ri[em]) != allp[em]:
            raise ValueError(f"{em}: 통·반 합과 인구 및 세대현황 총인구가 다르다")
    return Population(
        reference_date=ref_dates.pop(),
        ri=ri,
        emd_total={em: allp[em] for em in config.EMD_NAMES},
        emd_65plus={em: old[em] for em in config.EMD_NAMES},
        city_total=allp["공주시"],
        city_65plus=old["공주시"],
        dong_total={d: allp[d] for d in config.DONGS},
    )


def to_facts(p: Population) -> dict:
    """계산에 쓰는 값만(원본 표 전체가 아니다)."""
    return {
        "reference_date": p.reference_date,
        "emds": [
            {
                "name": em,
                "total": p.emd_total[em],
                "pop65": p.emd_65plus[em],
                "ri": [[name, n] for name, n in p.ri[em]],
            }
            for em in config.EMD_NAMES
        ],
        "city_total": p.city_total,
        "city_65plus": p.city_65plus,
        "dong_total": p.dong_total,
    }


def from_facts(f: dict) -> Population:
    return Population(
        reference_date=f["reference_date"],
        ri={e["name"]: [(name, int(n)) for name, n in e["ri"]] for e in f["emds"]},
        emd_total={e["name"]: e["total"] for e in f["emds"]},
        emd_65plus={e["name"]: e["pop65"] for e in f["emds"]},
        city_total=f["city_total"],
        city_65plus=f["city_65plus"],
        dong_total=f["dong_total"],
    )
